# %%
import pandas as pd
import numpy as np
import statsmodels.formula.api as smf

from strategic_plans.library import start
from statsmodels.stats.multitest import multipletests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


NORMALIZED = True
P_ADJUSTMENT = "fdr_bh"


# ---------------- Helpers ----------------

def significance_stars(p):
    """Return stars based on an adjusted p-value."""
    if pd.isna(p):
        return ""
    if p < 0.001:
        return "***"
    if p < 0.01:
        return "**"
    if p < 0.05:
        return "*"
    return ""


def topic_label(topic_code, top_topics_df):
    """Convert a modeled topic variable to its display label."""
    base = topic_code.removesuffix("_norm")

    match = top_topics_df.loc[
        top_topics_df["Topic ID"].astype(str) == str(base),
        "Topic Code"
    ]

    if match.empty:
        return str(base)

    return match.iloc[0]


def fit_contrast(data, outcome, group_variable, reference, comparison,
                 include_state_fixed_effects=True):
    """
    Estimate comparison minus reference.

    Returns the contrast, standard error, unadjusted p-value,
    confidence interval, model N, and the reference group's raw mean.
    """
    group_term = (
        f"C({group_variable}, Treatment(reference='{reference}'))"
    )

    controls = [
        "improvement_plan",
        "form_plan",
        "word_count",
    ]

    if include_state_fixed_effects:
        controls.insert(0, "C(state)")

    formula = f"{outcome} ~ {group_term} + " + " + ".join(controls)

    model = smf.ols(
        formula,
        data=data,
        missing="drop"
    ).fit(cov_type="HC1")

    parameter = f"{group_term}[T.{comparison}]"

    estimate = model.params.get(parameter, np.nan)
    standard_error = model.bse.get(parameter, np.nan)
    p_value = model.pvalues.get(parameter, np.nan)

    if parameter in model.params.index:
        confidence_interval = model.conf_int().loc[parameter]
        ci_low = confidence_interval.iloc[0]
        ci_high = confidence_interval.iloc[1]
    else:
        ci_low = np.nan
        ci_high = np.nan

    reference_mean = data.loc[
        data[group_variable] == reference,
        outcome
    ].mean()

    return {
        "estimate": estimate,
        "standard_error": standard_error,
        "p_value": p_value,
        "ci_low": ci_low,
        "ci_high": ci_high,
        "n": int(model.nobs),
        "reference_mean": reference_mean,
    }


# ---------------- Data ----------------

top_topics_df = pd.read_excel(
    start.RESULTS_DIR + "top_topics.xlsx"
)

merge_df = pd.read_excel(
    start.DATA_DIR + "clean/sample_inclusion_with_topics_and_codes.xlsx"
)

merge_df = merge_df.loc[merge_df["state"] != "PA"].copy()


# ---------------- Topic ordering ----------------

suffix = "_norm" if NORMALIZED else ""

top_topics_df = top_topics_df.sort_values(
    ["Group Order", "Average Prevalence"],
    ascending=[True, False]
)

TOPIC_GROUPS = []

for group_name, group_rows in top_topics_df.groupby(
    "Parent Code",
    sort=False
):
    members = [
        f"{topic_id}{suffix}"
        for topic_id in group_rows["Topic ID"].astype(str)
    ]

    TOPIC_GROUPS.append((group_name, members))

ordered_topics = [
    topic
    for _, topics in TOPIC_GROUPS
    for topic in topics
]


# ---------------- Comparison variables ----------------

# Urbanicity: Urban is the reference.
merge_df["urbanicity"] = np.select(
    [
        merge_df["urban"].eq(1),
        merge_df["suburb"].eq(1),
        merge_df["town"].eq(1),
    ],
    [
        "urban",
        "suburb",
        "town",
    ],
    default="rural",
)

# Region: Northeast is the reference.
merge_df["region_cat"] = np.select(
    [
        merge_df["midwest"].eq(1),
        merge_df["west"].eq(1),
        merge_df["south"].eq(1),
    ],
    [
        "midwest",
        "west",
        "south",
    ],
    default="northeast",
)

# Racial-demographic quartiles.
# Scaling to 0–100 is unnecessary because qcut depends only on rank.
merge_df["demographic_quartile"] = pd.qcut(
    merge_df["percent_race_black_hispanic"],
    q=4,
    labels=["Q1", "Q2", "Q3", "Q4"],
)

# Political quartiles.
merge_df["politics_quartile"] = pd.qcut(
    merge_df["district_pct_trump"],
    q=4,
    labels=["Q1", "Q2", "Q3", "Q4"],
)


# ---------------- Contrast specifications ----------------

CONTRASTS = {
    "Rural − Urban": {
        "group_variable": "urbanicity",
        "reference": "urban",
        "comparison": "rural",
        "include_state_fixed_effects": True,
    },
    "South − Northeast": {
        "group_variable": "region_cat",
        "reference": "northeast",
        "comparison": "south",

        # Region is perfectly collinear with state fixed effects.
        "include_state_fixed_effects": False,
    },
    "Highest − Lowest Black/Hispanic": {
        "group_variable": "demographic_quartile",
        "reference": "Q1",
        "comparison": "Q4",
        "include_state_fixed_effects": True,
    },
    "Most − Least Republican": {
        "group_variable": "politics_quartile",
        "reference": "Q1",
        "comparison": "Q4",
        "include_state_fixed_effects": True,
    },
}


# ---------------- Estimate contrasts ----------------

records = []

for topic_code in ordered_topics:
    topic_name = topic_label(topic_code, top_topics_df)

    for contrast_name, specification in CONTRASTS.items():
        result = fit_contrast(
            data=merge_df,
            outcome=topic_code,
            group_variable=specification["group_variable"],
            reference=specification["reference"],
            comparison=specification["comparison"],
            include_state_fixed_effects=(
                specification["include_state_fixed_effects"]
            ),
        )

        records.append({
            "topic_code": topic_code,
            "topic_name": topic_name,
            "contrast": contrast_name,
            **result,
        })

results = pd.DataFrame(records)


# ---------------- Multiple-testing adjustment ----------------

# Treat the 13 tests for each district characteristic as one family.
results["adjusted_p_value"] = np.nan

for contrast_name in CONTRASTS:
    mask = (
        results["contrast"].eq(contrast_name)
        & results["p_value"].notna()
    )

    results.loc[mask, "adjusted_p_value"] = multipletests(
        results.loc[mask, "p_value"],
        method=P_ADJUSTMENT,
    )[1]


# Convert proportions to percentage-point differences for presentation.
results["estimate_pp"] = results["estimate"] * 100
results["standard_error_pp"] = results["standard_error"] * 100
results["ci_low_pp"] = results["ci_low"] * 100
results["ci_high_pp"] = results["ci_high"] * 100
results["reference_mean_pp"] = results["reference_mean"] * 100


# ---------------- Workbook ----------------

wb = Workbook()
ws = wb.active
ws.title = "Variation Contrasts"

headers = ["Topic"] + list(CONTRASTS.keys())

for column, header in enumerate(headers, start=1):
    cell = ws.cell(row=1, column=column, value=header)
    cell.font = Font(bold=True)
    cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True,
    )

ws.freeze_panes = "B2"

current_row = 2

for group_name, topics in TOPIC_GROUPS:
    # Group heading
    ws.cell(
        row=current_row,
        column=1,
        value=group_name,
    ).font = Font(bold=True, italic=True)

    current_row += 1

    for topic_code in topics:
        topic_rows = results.loc[
            results["topic_code"].eq(topic_code)
        ]

        if topic_rows.empty:
            continue

        topic_name = topic_rows["topic_name"].iloc[0]

        # Estimate row
        ws.cell(
            row=current_row,
            column=1,
            value=topic_name,
        )

        # Standard-error row
        ws.cell(
            row=current_row + 1,
            column=1,
            value="",
        )

        for column, contrast_name in enumerate(
            CONTRASTS.keys(),
            start=2
        ):
            result = topic_rows.loc[
                topic_rows["contrast"].eq(contrast_name)
            ].iloc[0]

            estimate = result["estimate_pp"]
            standard_error = result["standard_error_pp"]
            adjusted_p = result["adjusted_p_value"]

            stars = significance_stars(adjusted_p)

            estimate_cell = ws.cell(
                row=current_row,
                column=column,
            )

            se_cell = ws.cell(
                row=current_row + 1,
                column=column,
            )

            if pd.notna(estimate):
                estimate_cell.value = f"{estimate:.1f}{stars}"
                se_cell.value = f"({standard_error:.1f})"
            else:
                estimate_cell.value = "—"
                se_cell.value = ""

            estimate_cell.alignment = Alignment(
                horizontal="center"
            )

            se_cell.alignment = Alignment(
                horizontal="center"
            )

            se_cell.font = Font(
                italic=True,
                color="666666",
            )

        current_row += 2


# ---------------- Notes ----------------

current_row += 1

note = (
    "Note. Entries are adjusted differences in normalized topic prevalence, "
    "reported in percentage points; heteroskedasticity-robust HC1 standard "
    "errors appear in parentheses. Positive estimates indicate greater "
    "prevalence in the first group named in the column heading. Urbanicity, "
    "racial-demographic, and political models include state fixed effects, "
    "improvement-plan status, standardized-plan status, and plan word count. "
    "The regional models omit state fixed effects because region is determined "
    "by state. Stars are based on Benjamini–Hochberg-adjusted p-values within "
    "each column: * p < .05, ** p < .01, *** p < .001."
)

ws.cell(
    row=current_row,
    column=1,
    value=note,
)

ws.merge_cells(
    start_row=current_row,
    start_column=1,
    end_row=current_row,
    end_column=len(headers),
)

ws.cell(
    row=current_row,
    column=1,
).alignment = Alignment(
    wrap_text=True,
    vertical="top",
)

ws.row_dimensions[current_row].height = 75


# ---------------- Formatting ----------------

column_widths = {
    1: 38,
    2: 17,
    3: 20,
    4: 27,
    5: 22,
}

for column, width in column_widths.items():
    ws.column_dimensions[
        get_column_letter(column)
    ].width = width

ws.row_dimensions[1].height = 42


# ---------------- Export supporting results ----------------

details_ws = wb.create_sheet("Underlying Estimates")

detail_columns = [
    "topic_name",
    "contrast",
    "reference_mean_pp",
    "estimate_pp",
    "standard_error_pp",
    "ci_low_pp",
    "ci_high_pp",
    "p_value",
    "adjusted_p_value",
    "n",
]

for column, variable in enumerate(detail_columns, start=1):
    details_ws.cell(
        row=1,
        column=column,
        value=variable,
    ).font = Font(bold=True)

for row_number, record in enumerate(
    results[detail_columns].itertuples(index=False),
    start=2,
):
    for column, value in enumerate(record, start=1):
        details_ws.cell(
            row=row_number,
            column=column,
            value=value,
        )


# ---------------- Save ----------------

output_path = (
    start.RESULTS_DIR
    + "topic_variation_endpoint_contrasts.xlsx"
)

wb.save(output_path)

print(f"Table exported to {output_path}")