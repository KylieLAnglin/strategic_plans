# %%
import pandas as pd

from strategic_plans.library import start, format_tables
import statsmodels.api as sm
import statsmodels.formula.api as smf
from openpyxl import load_workbook
import os
from openpyxl import Workbook

RESULTS_FILE = start.RESULTS_DIR + "Three Samples Characteristics.xlsx"

# %%
code_df = pd.read_csv(start.DATA_DIR + "clean/plans_codes_characteristics.csv")
sample_df = pd.read_csv(start.DATA_DIR + "clean/stratified_sample_characteristics.csv")
plan_df = pd.read_csv(start.DATA_DIR + "/clean/meta_data_df.csv")

df = sample_df.merge(plan_df[["leaid"]], on="leaid", how="left", indicator="_merge_plan") # 617 with plan. 414 without.
df = df.merge(code_df[["leaid"]], on="leaid", how="left", indicator="_merge_code") # 108 handcoded, 509 machine coded
df = df.drop_duplicates(subset=["leaid"], keep="first")
# %%
df["temp"] = df.northeast + df.midwest + df.south + df.west
# %%
binary_characteristics = [
    "northeast",
    "midwest",
    "south",
    "west",
    "urban",
    "suburb",
    "town",
    "rural",
]
continuous_characteristics = [
    "enrollment_in_thousands",
    "percent_race_black_hispanic",
    "percent_race_other",
    "percent_race_white",
    "percent_frl",
    "mean_test_score",
    "adults_w_ba",
    "district_pct_trump",
    "competitive",
]


file = RESULTS_FILE

if not os.path.exists(file):
    wb = Workbook()
    ws = wb.active
    ws.title = "descriptives"
    wb.save(file)

wb = load_workbook(file)
ws = wb["descriptives"]


# %%
# Column 2 = "Sample"
row = 2
col = 2
for char in binary_characteristics:
    mean_value = df[char].mean().round(2)
    ws.cell(row=row, column=col).value = mean_value
    row = row + 1
    sd_value = round(df[char].std(), 2)
    ws.cell(row=row, column=col).value = "[" + str(sd_value) + "]"
    row = row + 1

wb.save(file)

for char in continuous_characteristics:
    mean_value = df[char].mean().round(2)
    ws.cell(row=row, column=col).value = mean_value
    row = row + 1
    sd_value = round(df[char].std(), 2)
    ws.cell(row=row, column=col).value = "[" + str(sd_value) + "]"
    row = row + 1

ws.cell(row=row, column=col).value = len(df)


wb.save(file)

# Column 3 = "Sample with Plan"
row = 2
col = 3
df_plan = df[df._merge_plan == "both"]
for char in binary_characteristics:
    mean_value = df_plan[char].mean().round(2)
    ws.cell(row=row, column=col).value = mean_value
    row = row + 1
    sd_value = round(df_plan[char].std(), 2)
    ws.cell(row=row, column=col).value = "[" + str(sd_value) + "]"
    row = row + 1

for char in continuous_characteristics:
    mean_value = df_plan[char].mean().round(2)
    ws.cell(row=row, column=col).value = mean_value
    row = row + 1
    sd_value = round(df_plan[char].std(), 2)
    ws.cell(row=row, column=col).value = "[" + str(sd_value) + "]"
    row = row + 1
ws.cell(row=row, column=col).value = len(df_plan)

wb.save(file)

# Column 4 = "Sample with Codes"
row = 2
col = 4
df_code = df[df._merge_code == "both"]
for char in binary_characteristics:
    mean_value = df_code[char].mean().round(2)
    ws.cell(row=row, column=col).value = mean_value
    row = row + 1
    sd_value = round(df_code[char].std(), 2)
    ws.cell(row=row, column=col).value = "[" + str(sd_value) + "]"
    row = row + 1

for char in continuous_characteristics:
    mean_value = df_code[char].mean().round(2)
    ws.cell(row=row, column=col).value = mean_value
    row = row + 1
    sd_value = round(df_code[char].std(), 2)
    ws.cell(row=row, column=col).value = "[" + str(sd_value) + "]"
    row = row + 1
ws.cell(row=row, column=col).value = len(df_code)

# %%
wb.save(file)
# %%
# Print districts with highest enrollment and whether they are in code_df and plan_df, including district name
top_enrollment = df.nlargest(10, "enrollment_in_thousands")[["leaid", "enrollment_in_thousands", "_merge_code", "_merge_plan"]]
top_enrollment["is_in_code_df"] = top_enrollment["_merge_code"] == "both"
top_enrollment["is_in_plan_df"] = top_enrollment["_merge_plan"] == "both"
# Merge with plan_df to get district name
top_enrollment = top_enrollment.merge(plan_df[["leaid", "district"]], on="leaid", how="left")
print("Top 10 districts by enrollment:")
print(top_enrollment[["leaid", "district", "enrollment_in_thousands", "is_in_code_df", "is_in_plan_df"]])
# %%
