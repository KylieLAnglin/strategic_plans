# %%
"""
  - Purpose: Hierarchical code prevalence table across plans
  - Key operations:
    - One row per codebook node (every level), ordered as a tree: each
      node followed by its children, siblings sorted by prevalence
    - Prevalence uses each node's level aggregate column from
      plans_codes.csv (the node's own code or anything nested under it)
    - Excludes the Student Subgroups family
    - Outputs: results/hierarchical_prevalence.xlsx (APA-formatted sheets
      'Table 1' = goals for students and 'Table 2' = goals for the district
      and community, each with an indented stub column; sheet 'data' is raw)
"""
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from strategic_plans.library import start

EXCLUDED_FAMILY = "Student subgroups"

# %% Load data and the codebook export
df = pd.read_csv(start.DATA_DIR + "clean/plans_codes.csv")
codebook_df = pd.read_csv(start.LATEST_CODEBOOK)
number_plans = df.leaid.nunique()

# %% Select the nodes for the table
# Drop the excluded family (the node itself and everything nested under it)
level_title_columns = [col for col in codebook_df.columns if col.startswith("level") and col.endswith("_title")]
in_excluded_family = codebook_df[level_title_columns].fillna("").eq(EXCLUDED_FAMILY).any(axis=1)
table_codebook = codebook_df[~in_excluded_family].copy()

# Each node is counted through its own level aggregate column
table_codebook["node_column"] = table_codebook.apply(
    lambda row: row[f"level{int(row.level)}_column"], axis=1
)

# Nodes without a data column (codes dropped from the dataset, categories
# with no coded members) are excluded
missing_column = ~table_codebook.node_column.isin(df.columns)
for _, dropped_row in table_codebook[missing_column].iterrows():
    print(f"Skipping node with no data column: {dropped_row.code_title}")
table_codebook = table_codebook[~missing_column]

print(f"{number_plans} plans, {len(table_codebook)} codebook nodes in table")

# %% Prevalence per node
for node_column in table_codebook.node_column:
    df[node_column] = pd.to_numeric(df[node_column], errors="coerce").fillna(0)

table_codebook["count_plans"] = table_codebook.node_column.map(
    df[table_codebook.node_column].sum()
)
table_codebook["proportion"] = table_codebook.count_plans / number_plans

# %% Order rows as a tree: parents first, siblings by prevalence
# Sort keys per level: the count of the row's ancestor at that level (so
# subtrees stay together, highest-prevalence first), with blank levels
# sorting ahead of children (so each parent row leads its subtree)
ancestor_counts = {
    (node_row.level, node_row.code_title): node_row.count_plans
    for _, node_row in table_codebook.iterrows()
}
max_level = int(table_codebook.level.max())
sort_columns = []
for level_number in range(1, max_level + 1):
    ancestor_titles = table_codebook[f"level{level_number}_title"].fillna("")
    table_codebook[f"sort{level_number}_count"] = [
        -ancestor_counts.get((level_number, title), np.inf) if title else -np.inf
        for title in ancestor_titles
    ]
    table_codebook[f"sort{level_number}_title"] = ancestor_titles
    sort_columns += [f"sort{level_number}_count", f"sort{level_number}_title"]
table_codebook = table_codebook.sort_values(by=sort_columns)

# %% Raw data sheet: one column per level holding the node's title
hierarchical_table = pd.DataFrame(index=table_codebook.index)
for level_number in range(1, max_level + 1):
    hierarchical_table[f"level{level_number}"] = np.where(
        table_codebook.level == level_number, table_codebook.code_title, ""
    )
hierarchical_table["count_plans"] = table_codebook.count_plans.astype(int)
hierarchical_table["proportion"] = table_codebook.proportion
hierarchical_table["variable"] = table_codebook.node_column

# %% Build the APA-formatted sheets
# Table 1: goals for students (the Students subtree, with its level-2
# branches flush left). Table 2: goals for the district and community (all
# other level-1 branches). APA table style: three horizontal rules only
# (above and below the header row, below the last body row), no vertical
# lines, indented stub column for the hierarchy, proportions without a
# leading zero
apa_base_font = Font(name="Times New Roman", size=12)
apa_italic_font = Font(name="Times New Roman", size=12, italic=True)
apa_bold_font = Font(name="Times New Roman", size=12, bold=True)
rule_below = Border(bottom=Side(style="thin"))
rule_above_and_below = Border(top=Side(style="thin"), bottom=Side(style="thin"))

students_count = int(
    table_codebook.loc[table_codebook.code_title == "Students", "count_plans"].iloc[0]
)
students_codebook = table_codebook[
    (table_codebook.level1_title == "Students") & (table_codebook.level > 1)
]
district_codebook = table_codebook[table_codebook.level1_title != "Students"]

prevalence_note = (
    f"Note. N = {number_plans} district strategic plans. Prevalence for a "
    "higher-level goal reflects plans in which that goal or any goal nested "
    "under it was applied."
)
table_specs = [
    (
        "Table 1",
        "Prevalence of Goals for Students in District Strategic Plans",
        students_codebook,
        2,  # level-2 branches sit flush left
        prevalence_note
        + f" {students_count} of {number_plans} plans included at least one "
        "goal for students. Student subgroup codes are excluded.",
    ),
    (
        "Table 2",
        "Prevalence of Goals for the District and Community in District Strategic Plans",
        district_codebook,
        1,
        prevalence_note,
    ),
]

workbook = Workbook()
workbook.remove(workbook.active)

for sheet_name, table_title, sheet_codebook, flush_level, table_note in table_specs:
    apa_sheet = workbook.create_sheet(sheet_name)
    apa_sheet["A1"] = sheet_name
    apa_sheet["A1"].font = apa_bold_font
    apa_sheet["A2"] = table_title
    apa_sheet["A2"].font = apa_italic_font

    header_row = 4
    apa_sheet.cell(row=header_row, column=1, value="Goal")
    apa_sheet.cell(row=header_row, column=2, value="n")
    apa_sheet.cell(row=header_row, column=3, value="Proportion")
    for column_number in [1, 2, 3]:
        header_cell = apa_sheet.cell(row=header_row, column=column_number)
        header_cell.font = apa_italic_font if column_number == 2 else apa_base_font
        header_cell.border = rule_above_and_below
        header_cell.alignment = Alignment(horizontal="left" if column_number == 1 else "center")

    body_start_row = header_row + 1
    for row_offset, (_, node_row) in enumerate(sheet_codebook.iterrows()):
        excel_row = body_start_row + row_offset
        name_cell = apa_sheet.cell(row=excel_row, column=1, value=node_row.code_title)
        name_cell.font = apa_base_font
        name_cell.alignment = Alignment(
            horizontal="left", indent=int(node_row.level) - flush_level
        )
        count_cell = apa_sheet.cell(row=excel_row, column=2, value=int(node_row.count_plans))
        count_cell.font = apa_base_font
        count_cell.alignment = Alignment(horizontal="center")
        proportion_cell = apa_sheet.cell(row=excel_row, column=3, value=float(node_row.proportion))
        proportion_cell.font = apa_base_font
        proportion_cell.alignment = Alignment(horizontal="center")
        proportion_cell.number_format = "#.00"  # APA: no leading zero on proportions

    last_body_row = body_start_row + len(sheet_codebook) - 1
    for column_number in [1, 2, 3]:
        apa_sheet.cell(row=last_body_row, column=column_number).border = rule_below

    note_cell = apa_sheet.cell(row=last_body_row + 2, column=1)
    note_cell.value = table_note
    note_cell.font = apa_base_font
    note_cell.alignment = Alignment(horizontal="left", wrap_text=True)

    apa_sheet.column_dimensions["A"].width = 58
    apa_sheet.column_dimensions["B"].width = 8
    apa_sheet.column_dimensions["C"].width = 12

# %% Export (APA sheet plus a raw data sheet)
data_sheet = workbook.create_sheet("data")
for data_row in dataframe_to_rows(hierarchical_table, index=False, header=True):
    data_sheet.append(data_row)

workbook.save(start.RESULTS_DIR + "hierarchical_prevalence.xlsx")
print(f"Saved to {start.RESULTS_DIR}hierarchical_prevalence.xlsx")
print(hierarchical_table.drop(columns="variable").to_string(index=False))

# %%
