# %%
import pandas as pd
import numpy as np
from strategic_plans.library import start
import statsmodels.formula.api as smf
from statsmodels.stats.multitest import multipletests
from openpyxl import Workbook
from openpyxl.styles import Font

NORMALIZED = True

# ---------------- Helpers ----------------
def starify(p):
    return f"{p:.2f}" + ("***" if p < 0.001 else "**" if p < 0.01 else "*" if p < 0.05 else "")

def topic_label(topic_code, top_topics_df):
    base = topic_code.replace("_norm", "")
    return top_topics_df.loc[top_topics_df['Topic ID'] == base, 'Topic Code'].iloc[0]

# ---------------- Data ----------------
top_topics_df = pd.read_excel(start.RESULTS_DIR + 'top_topics.xlsx')
merge_df = pd.read_excel(start.DATA_DIR + "clean/sample_inclusion_with_topics_and_codes.xlsx")
merge_df = merge_df[merge_df.state != "PA"]

# Grouping, names, inclusion, and merges are all chosen in the ContentCoder app
# and flow here through top_topics.xlsx (see Table5b). Groups follow the app's
# LDA-tab order (Group Order); topics within a group are ordered by prevalence.
suffix = "_norm" if NORMALIZED else ""
top_topics_df = top_topics_df.sort_values(
    ["Group Order", "Average Prevalence"], ascending=[True, False]
)
TOPIC_GROUPS = []
for group_name, group_rows in top_topics_df.groupby("Parent Code", sort=False):
    members = [f"{topic_id}{suffix}" for topic_id in group_rows["Topic ID"]]
    TOPIC_GROUPS.append((group_name, members))

ordered_topics = [topic for _, topics in TOPIC_GROUPS for topic in topics]

# ---------------- Workbook ----------------
wb = Workbook()
ws = wb.active
ws.title = "Urbanicity Coefficients (State FE)"
headers = ["Topic", "Urban Mean", "Suburb", "Town", "Rural", "Unadj P-value", "Adj P-value"]
for c, h in enumerate(headers, 1):
    ws.cell(row=1, column=c, value=h).font = Font(bold=True)

# ---------------- Regressions ----------------
merge_df['urbanicity'] = np.where(merge_df["urban"] == 1, "urban",
                                  np.where(merge_df["suburb"] == 1, "suburb",
                                           np.where(merge_df["town"] == 1, "town", "rural")))
urbanicity_levels = ["suburb", "town", "rural"]

p_values = []
topic_results = []

for topic_code in ordered_topics:
    topic_name = topic_label(topic_code, top_topics_df)

    formula = (
        f"{topic_code} ~ C(urbanicity, Treatment(reference='urban')) + "
        f"C(state) + improvement_plan + form_plan + word_count"
    )
    model = smf.ols(formula, data=merge_df, missing='drop').fit()

    urban_mean = merge_df.loc[merge_df['urbanicity'] == 'urban', topic_code].mean()

    coefficients = {lvl: model.params.get(f"C(urbanicity, Treatment(reference='urban'))[T.{lvl}]", np.nan)
                    for lvl in urbanicity_levels}
    std_errors = {lvl: model.bse.get(f"C(urbanicity, Treatment(reference='urban'))[T.{lvl}]", np.nan)
                  for lvl in urbanicity_levels}

    f_test = model.f_test(
        "C(urbanicity, Treatment(reference='urban'))[T.suburb] = 0, "
        "C(urbanicity, Treatment(reference='urban'))[T.town] = 0, "
        "C(urbanicity, Treatment(reference='urban'))[T.rural] = 0"
    )
    p_value = float(f_test.pvalue)

    p_values.append(p_value)
    topic_results.append((topic_name, urban_mean, coefficients, std_errors, p_value))

# ---------------- Adjust & Write ----------------
_, adj_p_values, _, _ = multipletests(p_values, method='fdr_bh')

row = 2
for i, (topic_name, urban_mean, coefficients, std_errors, p_value) in enumerate(topic_results):
    ws.cell(row=row, column=1, value=topic_name)
    ws.cell(row=row, column=2, value=f"{urban_mean:.2f}")
    ws.cell(row=row, column=3, value=f"{coefficients['suburb']:.2f}")
    ws.cell(row=row, column=4, value=f"{coefficients['town']:.2f}")
    ws.cell(row=row, column=5, value=f"{coefficients['rural']:.2f}")
    ws.cell(row=row, column=6, value=starify(p_value))
    ws.cell(row=row, column=7, value=starify(adj_p_values[i]))

    row += 1
    ws.cell(row=row, column=3, value=f"({std_errors['suburb']:.2f})")
    ws.cell(row=row, column=4, value=f"({std_errors['town']:.2f})")
    ws.cell(row=row, column=5, value=f"({std_errors['rural']:.2f})")
    row += 1

# ---------------- Group headers ----------------
current_row = 2
for group_name, topics in TOPIC_GROUPS:
    ws.insert_rows(current_row)
    ws.cell(row=current_row, column=1, value=group_name).font = Font(bold=True, italic=True)
    current_row += 1 + (len(topics) * 2)

# ---------------- Save ----------------
output_path = start.RESULTS_DIR + 'Table6_urbanicity_topics_coefficients.xlsx'
wb.save(output_path)
print(f"\nTable exported to {output_path}")