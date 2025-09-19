# %%
import pandas as pd
import numpy as np
from strategic_plans.library import start
import statsmodels.formula.api as smf
from statsmodels.stats.multitest import multipletests
from openpyxl import Workbook
from openpyxl.styles import Font

# ---------------- Config ----------------
NORMALIZED = True  # flip False for raw prevalence

# ---------------- Helpers ----------------
def starify(p):
    return f"{p:.2f}" + ("***" if p < 0.001 else "**" if p < 0.01 else "*" if p < 0.05 else "")

def topic_label(topic_code, top_df):
    base = topic_code.replace("_norm", "")
    return top_df.loc[top_df['Topic ID'] == base, 'Topic Code'].iloc[0]

# ---------------- Load ----------------
top_topics_df = pd.read_excel(start.RESULTS_DIR + 'top_topics.xlsx')
merge_df = pd.read_excel(start.DATA_DIR + "clean/sample_inclusion_with_topics_and_codes.xlsx")

print(f"Analyzing {len(top_topics_df)} topics")
print(f"Number of districts: {len(merge_df)}")
print(f"Number of states: {merge_df['state'].nunique()}")

# Scale to 0–100
merge_df["district_pct_trump"] = merge_df["district_pct_trump"] * 100

# ---------------- Politics quartiles ----------------
merge_df['politics_quartile'] = pd.qcut(
    merge_df['district_pct_trump'], q=4, labels=['Q1', 'Q2', 'Q3', 'Q4']
)

print("\nPolitical Quartiles (% Trump Vote):")
quartile_stats = merge_df.groupby('politics_quartile')['district_pct_trump'].agg(['min', 'max', 'mean', 'count'])
for q in ['Q1', 'Q2', 'Q3', 'Q4']:
    s = quartile_stats.loc[q]
    print(f"{q}: {s['min']:.1f}% - {s['max']:.1f}% (mean: {s['mean']:.1f}%, n={int(s['count'])})")

# ---------------- Topic sets ----------------
if NORMALIZED:
    merge_df['Academic_Achievement_norm'] = merge_df['Topic_9_norm'] + merge_df['Topic_16_norm']
    top_topics_df = pd.concat([
        top_topics_df,
        pd.DataFrame({'Topic ID': ['Academic_Achievement'], 'Topic Code': ['Academic Achievement']})
    ], ignore_index=True)
    TOPIC_GROUPS = [
        ("Academic Topics", ["Topic_1_norm", "Topic_8_norm", "Academic_Achievement_norm"]),
        ("Non-Academic Outcomes", ["Topic_3_norm", "Topic_5_norm", "Topic_12_norm", "Topic_14_norm"]),
        ("Family and Community", ["Topic_17_norm", "Topic_20_norm"]),
        ("Mechanisms", ["Topic_0_norm", "Topic_22_norm"])
    ]
else:
    merge_df['Academic_Achievement'] = merge_df['Topic_9'] + merge_df['Topic_16']
    top_topics_df = pd.concat([
        top_topics_df,
        pd.DataFrame({'Topic ID': ['Academic_Achievement'], 'Topic Code': ['Academic Achievement']})
    ], ignore_index=True)
    TOPIC_GROUPS = [
        ("Academic Topics", ["Topic_1", "Topic_8", "Academic_Achievement"]),
        ("Non-Academic Outcomes", ["Topic_3", "Topic_5", "Topic_12", "Topic_14"]),
        ("Family and Community", ["Topic_17", "Topic_20"]),
        ("Mechanisms", ["Topic_0", "Topic_22"])
    ]

ordered_topics = [t for _, topics in TOPIC_GROUPS for t in topics]
quartile_levels = ["Q2", "Q3", "Q4"]  # Q1 is reference

# ---------------- Workbook ----------------
wb = Workbook()
ws = wb.active
ws.title = "Politics Coefficients (State FE)"
headers = ["Topic", "Q1 Mean", "Q2", "Q3", "Q4", "Unadj P-value", "Adj P-value"]
for c, h in enumerate(headers, 1):
    ws.cell(row=1, column=c, value=h).font = Font(bold=True)

# ---------------- Regressions ----------------
p_values = []
topic_results = []

for topic_code in ordered_topics:
    topic_name = topic_label(topic_code, top_topics_df)

    formula = (
        f"{topic_code} ~ C(politics_quartile, Treatment(reference='Q1')) "
        f"+ C(state) + improvement_plan + form_plan + word_count"
    )
    model = smf.ols(formula, data=merge_df, missing='drop').fit()

    q1_mean = merge_df.loc[merge_df['politics_quartile'] == 'Q1', topic_code].mean()

    coefficients = {
        lvl: model.params.get(f"C(politics_quartile, Treatment(reference='Q1'))[T.{lvl}]", np.nan)
        for lvl in quartile_levels
    }
    std_errors = {
        lvl: model.bse.get(f"C(politics_quartile, Treatment(reference='Q1'))[T.{lvl}]", np.nan)
        for lvl in quartile_levels
    }

    # Explicit joint F-test
    f_test = model.f_test(
        "C(politics_quartile, Treatment(reference='Q1'))[T.Q2] = 0, "
        "C(politics_quartile, Treatment(reference='Q1'))[T.Q3] = 0, "
        "C(politics_quartile, Treatment(reference='Q1'))[T.Q4] = 0"
    )
    p_value = float(f_test.pvalue)

    p_values.append(p_value)
    topic_results.append((topic_name, q1_mean, coefficients, std_errors, p_value))

# ---------------- Adjust p-values ----------------
_, adj_p_values, _, _ = multipletests(p_values, method='fdr_bh')

# ---------------- Write to Excel ----------------
row = 2
for i, (topic_name, q1_mean, coefficients, std_errors, p_value) in enumerate(topic_results):
    ws.cell(row=row, column=1, value=topic_name)
    ws.cell(row=row, column=2, value=f"{q1_mean:.2f}")
    ws.cell(row=row, column=3, value=f"{coefficients['Q2']:.2f}")
    ws.cell(row=row, column=4, value=f"{coefficients['Q3']:.2f}")
    ws.cell(row=row, column=5, value=f"{coefficients['Q4']:.2f}")
    ws.cell(row=row, column=6, value=starify(p_value))
    ws.cell(row=row, column=7, value=starify(adj_p_values[i]))

    row += 1
    ws.cell(row=row, column=3, value=f"({std_errors['Q2']:.2f})")
    ws.cell(row=row, column=4, value=f"({std_errors['Q3']:.2f})")
    ws.cell(row=row, column=5, value=f"({std_errors['Q4']:.2f})")
    row += 1

# ---------------- Group separators ----------------
current_row = 2
for group_name, topics in TOPIC_GROUPS:
    ws.insert_rows(current_row)
    ws.cell(row=current_row, column=1, value=group_name).font = Font(bold=True, italic=True)
    current_row += 1 + (len(topics) * 2)

# ---------------- Notes & Save ----------------
ws.cell(row=current_row + 1, column=1, value="Note: Q1 (lowest % Trump vote) is the reference category").font = Font(italic=True)
if NORMALIZED:
    ws.cell(row=current_row + 2, column=1, value="Prevalence values normalized by sum of included topics").font = Font(italic=True)

output_path = start.RESULTS_DIR + 'Table11_politics_quartiles_coefficients.xlsx'
wb.save(output_path)
print(f"\nTable exported to {output_path}")