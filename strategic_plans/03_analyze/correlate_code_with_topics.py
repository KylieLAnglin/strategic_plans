# %%
import pandas as pd
import statsmodels.formula.api as smf
from openpyxl import Workbook
from openpyxl.styles import Font

from strategic_plans.library import start

# %%
# ------------------ SETUP ------------------
NORMALIZED = True
MIN_DISTRICTS = 30

INPUT_MERGE_DF = start.DATA_DIR + "clean/sample_inclusion_with_topics_and_codes.xlsx"
INPUT_TOP_TOPICS_DF = start.RESULTS_DIR + "top_topics.xlsx"
EXPORT_PATH = start.RESULTS_DIR + "code_topic_correlations.xlsx"


def starify(coefficient, p_value):
    stars = "***" if p_value < 0.001 else "**" if p_value < 0.01 else "*" if p_value < 0.05 else ""
    return f"{coefficient:.3f}{stars}"


def code_label(code_col):
    return code_col.replace("_applied", "").replace("code_", "").replace("_", " ").strip().capitalize()


# %%
# ------------------ LOAD DATA ------------------
merge_df = pd.read_excel(INPUT_MERGE_DF)
merge_df = merge_df[merge_df.state != "PA"]
merge_df = merge_df[merge_df.in_human_sample == 1]

top_topics_df = pd.read_excel(INPUT_TOP_TOPICS_DF)

suffix = "_norm" if NORMALIZED else ""
topic_cols = [f"{topic_id}{suffix}" for topic_id in top_topics_df["Topic ID"]]
topic_labels = list(top_topics_df["Topic Code"])

code_cols = [col for col in merge_df.columns if "applied" in col and "title" not in col]

# %%
# ------------------ FILTER CODES ------------------
code_district_counts = (merge_df[code_cols] == 1).sum()
included_codes = [col for col in code_cols if code_district_counts[col] >= MIN_DISTRICTS]

print(f"Codes total: {len(code_cols)}")
print(f"Codes with at least {MIN_DISTRICTS} districts: {len(included_codes)}")

# %%
# ------------------ WORKBOOK ------------------
wb = Workbook()
ws = wb.active
ws.title = "Code-Topic Coefficients"

ws.cell(row=1, column=1, value="Code").font = Font(bold=True)
for c, topic_label in enumerate(topic_labels, start=2):
    ws.cell(row=1, column=c, value=topic_label).font = Font(bold=True)

# %%
# ------------------ REGRESSIONS ------------------
row = 2
for code_col in included_codes:
    ws.cell(row=row, column=1, value=code_label(code_col))
    for c, topic_col in enumerate(topic_cols, start=2):
        model = smf.ols(f"{topic_col} ~ {code_col}", data=merge_df, missing="drop").fit()
        coefficient = model.params[code_col]
        std_error = model.bse[code_col]
        p_value = model.pvalues[code_col]

        ws.cell(row=row, column=c, value=starify(coefficient, p_value))
        ws.cell(row=row + 1, column=c, value=f"({std_error:.3f})")
    row += 2

# %%
# ------------------ EXPORT ------------------
wb.save(EXPORT_PATH)
print(f"Saved: {EXPORT_PATH}")
