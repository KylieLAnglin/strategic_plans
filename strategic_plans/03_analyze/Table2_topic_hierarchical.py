# %%
"""
  - Purpose: Hierarchical topic prevalence table, organized like the goal
    prevalence table (Table_hierarchical) with two levels: each LDA group
    row followed by its member topics, siblings sorted by prevalence
  - Columns match the topic prevalence table (Table5_top_topics): mean,
    25th, and 75th percentile of normalized prevalence, plus top words.
    Both are topic-level; a group row is a label only
  - A topic that merges several LDA topics gets one row per merged topic:
    prevalence sits on the first row, and each row carries one LDA topic's
    words, so the words justify combining them
  - Groups are ordered by group_order, the order set by hand in the app's
    LDA tab; groups sharing an order fall back to summed prevalence
  - Topic curation (names, groups, inclusion, merges) comes from the app's
    topics export (contentcoder/export_topics.py)
  - Outputs: results/topic_hierarchical_prevalence.xlsx (APA-formatted
    sheet 'Table 1' with an indented stub column; sheet 'data' is raw)
"""
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from strategic_plans.library import start

WORDS_IN_APA_SHEET = 5
WORDS_IN_DATA_SHEET = 10

# %% Load the topic curation and district-level prevalence
topic_naming_path = start.DATA_DIR + 'final_model/topic_naming_named.xlsx'
topic_naming_df = pd.read_excel(topic_naming_path)
topic_naming_df = topic_naming_df.rename(columns={'Unnamed: 0': 'topic_id'})
topic_naming_df = topic_naming_df.rename(
    columns={c: f'Word {i + 1}' for i, c in enumerate(range(WORDS_IN_DATA_SHEET))}
)
assert {'include_in_analysis', 'merged_into', 'group_order'} <= set(topic_naming_df.columns), (
    "topic_naming_named.xlsx predates the app's topics export; "
    "run contentcoder/export_topics.py to regenerate it"
)
topic_naming_df['topic_id_space'] = topic_naming_df['topic_id'].str.replace("_", " ", regex=False)

doc_topics_path = start.DATA_DIR + 'final_model/topic_model/doc_topics_grouped.xlsx'
doc_topics_df = pd.read_excel(doc_topics_path)
topic_cols = [str(i) for i in range(23)]
doc_topics_df = doc_topics_df.rename(columns={c: f"Topic {int(c)}" for c in topic_cols})

# %% Fold merged topics into their canonical topic (prevalences sum)
merged_topics_meta = topic_naming_df[
    topic_naming_df['merged_into'].notna() & (topic_naming_df['merged_into'] != '')
]
for _, merged_topic in merged_topics_meta.iterrows():
    target_col = merged_topic['merged_into'].replace("_", " ")
    source_col = merged_topic['topic_id_space']
    doc_topics_df[target_col] = doc_topics_df[target_col] + doc_topics_df[source_col]
    doc_topics_df = doc_topics_df.drop(columns=[source_col])

included_topics_meta = topic_naming_df[
    (topic_naming_df['include_in_analysis'] == 1)
    & ~topic_naming_df['topic_id'].isin(merged_topics_meta['topic_id'])
].copy()
included_topics_meta['parent_code'] = included_topics_meta['parent_code'].fillna("Ungrouped")
included_topic_cols = [
    t for t in included_topics_meta['topic_id_space'].tolist() if t in doc_topics_df.columns
]

# %% Row-normalize to proportion of included topics (same as Table5_top_topics)
included_values = doc_topics_df[included_topic_cols].apply(pd.to_numeric, errors='coerce')
row_sums = included_values.sum(axis=1)
normalized = included_values.div(row_sums.replace(0, np.nan), axis=0)
number_districts = int(normalized.notna().any(axis=1).sum())

# %% Prevalence statistics: one row per group, one per member topic
word_columns = [f'Word {i + 1}' for i in range(WORDS_IN_DATA_SHEET)]
table_rows = []
for group_name, group_meta in included_topics_meta.groupby('parent_code'):
    member_cols = [t for t in group_meta['topic_id_space'] if t in normalized.columns]
    group_order = group_meta['group_order'].iloc[0]
    # Only topics report statistics; group prevalence just breaks order ties
    group_series = normalized[member_cols].sum(axis=1, min_count=1).dropna()
    table_rows.append({
        'level': 1,
        'name': group_name,
        'sort_group_order': group_order,
        'sort_group_mean': group_series.mean(),
        'sort_topic_mean': np.nan,
        'sort_merge_rank': 0,
        'mean': np.nan,
        'p25': np.nan,
        'p75': np.nan,
        **{word: "" for word in word_columns},
    })
    for _, topic_meta in group_meta.iterrows():
        if topic_meta['topic_id_space'] not in normalized.columns:
            continue
        topic_series = normalized[topic_meta['topic_id_space']].dropna()
        table_rows.append({
            'level': 2,
            'name': topic_meta['code'],
            'sort_group_order': group_order,
            'sort_group_mean': group_series.mean(),
            'sort_topic_mean': topic_series.mean(),
            'sort_merge_rank': 0,
            'mean': topic_series.mean(),
            'p25': topic_series.quantile(0.25),
            'p75': topic_series.quantile(0.75),
            **{word: topic_meta[word] for word in word_columns},
        })
        # A merged-in LDA topic contributes only its words: the prevalence above
        # already includes it, and the words show what was combined
        merged_sources = merged_topics_meta[
            merged_topics_meta['merged_into'] == topic_meta['topic_id']
        ]
        for merge_rank, (_, source_meta) in enumerate(merged_sources.iterrows(), start=1):
            table_rows.append({
                'level': 2,
                'name': "",
                'sort_group_order': group_order,
                'sort_group_mean': group_series.mean(),
                'sort_topic_mean': topic_series.mean(),
                'sort_merge_rank': merge_rank,
                'mean': np.nan,
                'p25': np.nan,
                'p75': np.nan,
                **{word: source_meta[word] for word in word_columns},
            })

table_df = pd.DataFrame(table_rows)
# Groups in the order set in the app, each group row leading its topics (those
# by prevalence), with a topic's merged-in word rows following it
table_df = table_df.sort_values(
    by=['sort_group_order', 'sort_group_mean', 'level', 'sort_topic_mean', 'sort_merge_rank'],
    ascending=[True, False, True, False, True],
).drop(columns=['sort_group_order', 'sort_group_mean', 'sort_topic_mean', 'sort_merge_rank'])

print(f"{number_districts} districts, {len(included_topic_cols)} topics in "
      f"{included_topics_meta['parent_code'].nunique()} groups")

# %% Build the APA-formatted sheet (same style as Table_hierarchical)
apa_base_font = Font(name="Times New Roman", size=12)
apa_italic_font = Font(name="Times New Roman", size=12, italic=True)
apa_bold_font = Font(name="Times New Roman", size=12, bold=True)
rule_below = Border(bottom=Side(style="thin"))
rule_above_and_below = Border(top=Side(style="thin"), bottom=Side(style="thin"))

workbook = Workbook()
apa_sheet = workbook.active
apa_sheet.title = "Table 1"
apa_sheet["A1"] = "Table 1"
apa_sheet["A1"].font = apa_bold_font
apa_sheet["A2"] = "Prevalence of Topics in District Strategic Plans, by Topic Group"
apa_sheet["A2"].font = apa_italic_font

header_labels = ["Topic", "Mean", "25th", "75th"] + [
    f"Word {i + 1}" for i in range(WORDS_IN_APA_SHEET)
]
header_row = 4
for column_number, label in enumerate(header_labels, start=1):
    header_cell = apa_sheet.cell(row=header_row, column=column_number, value=label)
    header_cell.font = apa_base_font
    header_cell.border = rule_above_and_below
    header_cell.alignment = Alignment(horizontal="left" if column_number == 1 else "center")

body_start_row = header_row + 1
for row_offset, (_, table_row) in enumerate(table_df.iterrows()):
    excel_row = body_start_row + row_offset
    name_cell = apa_sheet.cell(row=excel_row, column=1, value=table_row['name'])
    name_cell.font = apa_base_font
    name_cell.alignment = Alignment(horizontal="left", indent=int(table_row['level']) - 1)
    for column_offset, stat in enumerate(['mean', 'p25', 'p75']):
        # Blank on a merged topic's extra word rows; the prevalence above covers them
        stat_value = None if pd.isna(table_row[stat]) else float(table_row[stat])
        stat_cell = apa_sheet.cell(row=excel_row, column=2 + column_offset, value=stat_value)
        stat_cell.font = apa_base_font
        stat_cell.alignment = Alignment(horizontal="center")
        stat_cell.number_format = "#.00"  # APA: no leading zero on proportions
    for word_offset in range(WORDS_IN_APA_SHEET):
        word_cell = apa_sheet.cell(
            row=excel_row, column=5 + word_offset,
            value=table_row[f'Word {word_offset + 1}'],
        )
        word_cell.font = apa_base_font
        word_cell.alignment = Alignment(horizontal="center")

last_body_row = body_start_row + len(table_df) - 1
for column_number in range(1, len(header_labels) + 1):
    apa_sheet.cell(row=last_body_row, column=column_number).border = rule_below

note_cell = apa_sheet.cell(row=last_body_row + 2, column=1)
note_cell.value = (
    f"Note. N = {number_districts} district strategic plans. Mean is the mean "
    "normalized prevalence across districts; 25th and 75th are the 25th and 75th "
    "percentile of normalized prevalence. Topics are ordered by prevalence within "
    "their group. Words are each topic's highest-probability terms. A "
    "topic that combines several LDA topics lists each one's words on its own row; "
    "the prevalence reported for the topic already includes all of them."
)
note_cell.font = apa_base_font
note_cell.alignment = Alignment(horizontal="left", wrap_text=True)

apa_sheet.column_dimensions["A"].width = 52
for column_letter in ["B", "C", "D"]:
    apa_sheet.column_dimensions[column_letter].width = 8
for word_offset in range(WORDS_IN_APA_SHEET):
    apa_sheet.column_dimensions[chr(ord("E") + word_offset)].width = 16

# %% Export (APA sheet plus a raw data sheet)
data_sheet = workbook.create_sheet("data")
for data_row in dataframe_to_rows(table_df, index=False, header=True):
    data_sheet.append(data_row)

output_path = start.RESULTS_DIR + "topic_hierarchical_prevalence.xlsx"
workbook.save(output_path)
print(f"Saved to {output_path}")
print(table_df[['level', 'name', 'mean', 'p25', 'p75', 'Word 1', 'Word 2']].round(3).to_string(index=False))

# %%
